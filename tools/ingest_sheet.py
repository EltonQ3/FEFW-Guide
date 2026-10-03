"""把共創表匯出的檔案轉成 source/sheet/*.json，並登記到 source/sheet/tabs.json。

支援：
- 騰訊文檔「導出為 → 本地 Excel」的 .xlsx（每個工作表一個分頁，合併儲存格會展開）
- 單一分頁的 .csv / .tsv（tools/fetch_sheet.py 存下的剪貼簿內容也是 TSV）

用法：
  python3 tools/ingest_sheet.py 共創表.xlsx
  python3 tools/ingest_sheet.py source/sheet/raw/*.tsv

已經登記過的分頁只更新資料，不覆蓋 tabs.json 裡手動調整的版面設定（標題、篩選欄、分組欄……）。
"""
from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SHEET = ROOT / "source" / "sheet"
TABS = SHEET / "tabs.json"

# 已知分頁名稱關鍵字 → 頁面代號與預設樣式。名稱以共創表分頁標題比對（包含即可）。
KNOWN = [
    ("兵种", {"slug": "classes", "page": "classes.html", "page_label": "兵种职业", "builtin": True}),
    ("外传", {"slug": "paralogues", "page": "paralogues.html", "page_label": "外传", "builtin": True}),
    ("血印", {"slug": "seals", "page": "seals.html", "page_label": "血印", "builtin": True}),
    ("加护", {"slug": "blessings", "page": "blessings.html", "page_label": "神之加护", "builtin": True}),
    ("特技", {"slug": "skills", "title": "特技", "en": "SKILLS", "icon": "classes", "color": "var(--azure)"}),
    ("战技", {"slug": "arts", "title": "战技", "en": "COMBAT ARTS", "icon": "classes", "color": "var(--crimson)"}),
    ("武器", {"slug": "weapons", "title": "武器", "en": "WEAPONS", "icon": "classes", "color": "var(--amber)"}),
    ("道具", {"slug": "items", "title": "道具", "en": "ITEMS", "icon": "seals", "color": "var(--jade)"}),
    ("军团", {"slug": "legion", "title": "军团兵", "en": "LEGIONS", "icon": "characters", "color": "var(--ember)"}),
    ("骑乘", {"slug": "mounts", "title": "骑乘", "en": "MOUNTS", "icon": "characters", "color": "var(--moon)"}),
    ("挖角", {"slug": "scouting", "title": "挖角名声", "en": "SCOUTING", "icon": "characters", "color": "var(--violet)"}),
    ("钓鱼", {"slug": "fishing", "title": "钓鱼", "en": "FISHING", "icon": "paralogues", "color": "var(--azure)"}),
    ("栽种", {"slug": "farming", "title": "栽种", "en": "FARMING", "icon": "paralogues", "color": "var(--jade)"}),
    ("支援", {"slug": "supports", "title": "支援", "en": "SUPPORTS", "icon": "characters", "color": "var(--crimson)"}),
    ("送礼", {"slug": "gifts", "title": "送礼", "en": "GIFTS", "icon": "characters", "color": "var(--amber)"}),
    ("茶会", {"slug": "teatime", "title": "茶会", "en": "TEA TIME", "icon": "blessings", "color": "var(--pearl)"}),
    ("喜爱", {"slug": "likes", "title": "角色喜爱", "en": "FAVORITES", "icon": "characters", "color": "var(--crimson)"}),
    ("公式", {"slug": "formulas", "title": "计算公式", "en": "FORMULAS", "icon": "about", "color": "var(--stone)"}),
]


def known_for(name: str) -> dict:
    for key, cfg in KNOWN:
        if key in name:
            return dict(cfg)
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "tab"
    return {"slug": f"tab-{slug}" if slug != "tab" else "tab", "title": re.sub(r"^\d+[-－_ ]*", "", name), "en": "FROM THE SHEET",
            "icon": "about", "color": "var(--gold)"}


def header_index(rows: list[list[str]]) -> int:
    widths = [sum(1 for v in r if v) for r in rows[:12]]
    if not widths:
        return 0
    need = max(2, int(max(widths) * 0.6))
    for i, w in enumerate(widths):
        if w >= need:
            return i
    return 0


def tidy(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    rows = [[("" if v is None else str(v)).strip() for v in r] for r in rows]
    h = header_index(rows)
    header = rows[h] if rows else []
    width = max([len(header)] + [len(r) for r in rows[h + 1:]]) if rows else 0
    keep = [i for i in range(width) if (i < len(header) and header[i]) or any(i < len(r) and r[i] for r in rows[h + 1:])]
    cols = [header[i] if i < len(header) and header[i] else f"栏{i + 1}" for i in keep]
    body = [[r[i] if i < len(r) else "" for i in keep] for r in rows[h + 1:]]
    body = [r for r in body if any(r)]
    return cols, body


def read_xlsx(path: Path) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        for rng in list(ws.merged_cells.ranges):
            v = ws.cell(rng.min_row, rng.min_col).value
            ws.unmerge_cells(str(rng))
            for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
                for cell in row:
                    cell.value = v
        out.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
    return out


def read_delimited(path: Path) -> list[tuple[str, list[list[str]]]]:
    text = path.read_text(encoding="utf-8-sig")
    delim = "\t" if path.suffix.lower() in (".tsv", ".txt") else ","
    return [(path.stem, list(csv.reader(text.splitlines(), delimiter=delim)))]


def main(paths: list[str]) -> None:
    tabs = json.loads(TABS.read_text(encoding="utf-8")) if TABS.exists() else []
    by_name = {t["name"]: t for t in tabs}
    for p in map(Path, paths):
        sheets = read_xlsx(p) if p.suffix.lower() in (".xlsx", ".xlsm") else read_delimited(p)
        for name, rows in sheets:
            cols, body = tidy(rows)
            cfg = by_name.get(name) or {"name": name, **known_for(name)}
            if not cfg.get("builtin"):
                cfg["file"] = f"{cfg['slug']}.json"
                cfg.setdefault("page", f"{cfg['slug']}.html")
                cfg.setdefault("page_label", cfg.get("title", name))
            cfg["status"] = f"已收录（{len(body)} 行）"
            (SHEET / f"{cfg['slug']}.json").write_text(json.dumps({"columns": cols, "rows": body}, ensure_ascii=False, indent=1) + "\n",
                                                       encoding="utf-8")
            if name not in by_name:
                tabs.append(cfg)
                by_name[name] = cfg
            print(f"{name}: {len(cols)} 栏 × {len(body)} 行 → source/sheet/{cfg['slug']}.json")
    TABS.write_text(json.dumps(tabs, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1:])
